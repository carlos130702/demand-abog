"""
Renderiza el JSON producido por el modelo en un .docx con el formato
tipico de una sentencia peruana (encabezado del juzgado, PARTE EXPOSITIVA,
PARTE CONSIDERATIVA numerada, PARTE RESOLUTIVA).

No depende de una plantilla con tags fragiles (docxtpl): construye el
documento parrafo por parrafo con python-docx, lo que permite controlar
negritas/subrayados exactamente como en un modelo real y evita romperse
si un campo viene vacio o mas largo de lo esperado.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(11)


def _style_run(run, bold=False, italic=False, underline=False):
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run.bold = bold
    run.italic = italic
    run.underline = underline
    return run


def _p(doc, text="", bold=False, italic=False, underline=False,
       align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8):
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    if text:
        _style_run(para.add_run(text), bold, italic, underline)
    return para


def _numbered_list(doc, items, start=1):
    for i, item in enumerate(items, start=start):
        _p(doc, f"{i}. {item}")


def build_docx(data: dict, output_path: str, nota_de_trabajo: str | None = None):
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Inches(11), Inches(8.5)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, m, Inches(1))

    enc = data.get("encabezado", {})

    if nota_de_trabajo:
        _p(doc, "NOTA DE TRABAJO — LEER ANTES DE USAR", bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        _p(doc, nota_de_trabajo, space_after=20)
        doc.add_page_break()

    # Encabezado
    _p(doc, enc.get("juzgado", "[JUZGADO]"), bold=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if enc.get("sede"):
        _p(doc, f"Sede: {enc['sede']}", italic=True,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

    campos = [
        ("EXPEDIENTE N°", enc.get("expediente", "[●]")),
        ("ESPECIALISTA", enc.get("especialista", "[●]")),
        ("MATERIA", enc.get("materia", "[●]")),
        ("DEMANDANTE", enc.get("demandante", "[●]")),
        ("DEMANDADO", enc.get("demandado", "[●]")),
    ]
    for etiqueta, valor in campos:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        _style_run(para.add_run(f"{etiqueta}\t: "), bold=True)
        _style_run(para.add_run(valor))

    _p(doc, "")
    _p(doc, "SENTENCIA N° [●]", bold=True, underline=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _p(doc, "RESOLUCIÓN N° [●]", bold=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _p(doc, "Lima, [●] de [●] del [●].",
       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)

    # I. PARTE EXPOSITIVA
    _p(doc, "I. PARTE EXPOSITIVA:", bold=True, underline=True,
       align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10)

    _p(doc, "A. Argumento de la demandante:", bold=True, underline=True)
    for parrafo in data.get("argumento_demandante", []):
        _p(doc, parrafo)

    _p(doc, "B. Argumento de la demandada:", bold=True, underline=True)
    for parrafo in data.get("argumento_demandado", []):
        _p(doc, parrafo)

    _p(doc, "C. Trámite del proceso:", bold=True, underline=True)
    for parrafo in data.get("tramite", []):
        _p(doc, parrafo)

    # II. PARTE CONSIDERATIVA
    _p(doc, "II. PARTE CONSIDERATIVA:", bold=True, underline=True, space_after=10)

    puntos = data.get("puntos_controvertidos", [])
    if puntos:
        _p(doc, "PUNTOS CONTROVERTIDOS:", bold=True, underline=True)
        _numbered_list(doc, puntos)

    ordinales = [
        "Primero", "Segundo", "Tercero", "Cuarto", "Quinto", "Sexto",
        "Séptimo", "Octavo", "Noveno", "Décimo", "Décimo Primero",
        "Décimo Segundo", "Décimo Tercero", "Décimo Cuarto", "Décimo Quinto",
    ]
    for i, considerando in enumerate(data.get("considerandos", [])):
        ordinal = ordinales[i] if i < len(ordinales) else f"Considerando {i + 1}"
        titulo = considerando.get("titulo", "").upper()
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(6)
        _style_run(para.add_run(f"{ordinal}: "), bold=True, underline=True)
        _style_run(para.add_run(titulo + ".-"), bold=True, underline=True)
        for parrafo in considerando.get("parrafos", []):
            _p(doc, parrafo)

    # III. PARTE RESOLUTIVA
    _p(doc, "III. PARTE RESOLUTIVA:", bold=True, underline=True, space_after=10)
    _p(doc, "RESUELVE:", bold=True, underline=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    _numbered_list(doc, data.get("fallo", []))

    notas = data.get("notas_para_revision", [])
    if notas:
        doc.add_page_break()
        _p(doc, "NOTAS PARA REVISIÓN (generadas automáticamente)", bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        _numbered_list(doc, notas)

    doc.save(output_path)
